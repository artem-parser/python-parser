from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from day7 import parse_book

url = 'https://books.toscrape.com'
books, count_five = parse_book(url)

wb = Workbook()
ws = wb.active
ws.title = 'Книги'

ws2 = wb.create_sheet('Статистика')

headers = ['id', 'title', 'price', 'rating', 'stock', 'url']
ws.append(headers)
ws2.append(['Всего книг', 'Книг с рейтингом Five'])
ws2.append([len(books), count_five])


header_font = Font(bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_align = Alignment(horizontal='center')

for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align

for cell in ws2[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align

for i, book in enumerate(books, start=1):
    ws.append([
    i,
    book['title'],
    book['price'],
    book['rating'],
    book['stock'],
    book['url']

])

for col in ws.columns:
    max_len = max(len(str(cell.value)) for cell in col if cell.value is not None)
    ws.column_dimensions[col[0].column_letter].width = max_len + 2

for col in ws2.columns:
    max_len = max(len(str(cell.value)) for cell in col if cell.value is not None)
    ws2.column_dimensions[col[0].column_letter].width = max_len + 2

wb.save('books.xlsx')
print(f'Готово! Сохранено {len(books)} книг в books.xlsx')