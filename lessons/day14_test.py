import requests
from bs4 import BeautifulSoup
import csv
import time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment



def get_page(url, headers):
    try:
        response = requests.get(url, headers=headers, timeout=10)
        response.encoding = 'utf-8'
        if response.status_code == 200:
            return response.text
        print(f'Статус-код {url} - {response.status_code}')
        return None
    except Exception as e:
        print(f'Ошибка {e} в {url}')
        return None

def parse_books(text):
    soup = BeautifulSoup(text, 'html.parser')
    books = soup.select('article.product_pod')
    all_books = []

    for i, book in enumerate(books, start=1):
        link = book.select_one('h3 a')
        price = book.select_one('p.price_color').get_text(strip=True)
        rating_text = book.select_one('p.star-rating')
        rating = rating_text['class'][1]
        stock = book.select_one('p.instock.availability').get_text(strip=True)
        href = link.get('href')
        title = link.get('title')
        all_books.append({
            'id': i,
            'title': title,
            'price': price,
            'stock': stock,
            'rating': rating,
            'href': href
        })
    return all_books


def save_to_csv(books, filename):
    with open(filename, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=['id', 'title', 'price', 'stock', 'rating', 'href'])
        writer.writeheader()
        for i, book in enumerate(books, start=1):
            book['id'] = i
            writer.writerow(book)
    print(f'Сохранено csv в {filename}')


def save_to_xlsx(books, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Книги'
    ws.append(['id', 'title', 'price', 'stock', 'rating', 'href'])

    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill(start_color='1faee9', end_color='1faee9', fill_type='solid')
    head_align = Alignment(horizontal='center')

    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = head_align

    for i, book in enumerate(books, start=1):
        ws.append([i, book['title'],book['price'], book['stock'],book['rating'],book['href']])

    for col in ws.columns:
        max_len = max(len(str(cell.value)) for cell in col if cell.value is not None)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2

    wb.save(filename)
    print(f'Сохранено xlsx в {filename}')


def main():
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 (KHTML, like Gecko) "
                      "Chrome/120.0.0.0 Safari/537.36"
    }
    url_base= 'https://books.toscrape.com/catalogue/category/books/young-adult_21/page-{}.html'
    url_count = 3
    all_books = []

    for i in range(1, url_count + 1):
        print(f'Парсим {i} страницу...')
        url = url_base.format(i)
        html = get_page(url, headers)
        if html:
            books = parse_books(html)
            all_books.extend(books)
            print(f'  Найдено: {len(books)} книг')
        time.sleep(2)

    print(f'Всего книг: {len(all_books)}')

    save_to_csv(all_books, 'dook_test.csv')
    save_to_xlsx(all_books, 'dook_test.xlsx')

    five_count = sum(1 for i in all_books if i['rating'] == 'Five')
    one_count = sum(1 for i in all_books if i['rating'] == 'One')
    prices = [float(b['price'].replace('£', '')) for b in all_books]
    price_avg = sum(prices) / len(all_books)

    print(f'Всего с рейтингом Five: {five_count}')
    print(f'Всего с рейтингом One: {one_count}')
    print(f'Средняя цена: {price_avg:.2f}')

if __name__ == '__main__':
    main()