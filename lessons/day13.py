import csv

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import time


def get_page(url, headers):
    try:
        response = requests.get(url, headers=headers, timeout=10)
        response.encoding = 'utf-8'
        if response.status_code == 200:
            return response.text
        else:
            print(f'Статус: {response.status_code} для {url}')
            return None
    except Exception as e:
        print(f'Ошибка при скачивании {url}: {e}')
        return None


def parse_page(html):
    soup = BeautifulSoup(html, 'html.parser')
    books = soup.select('article.product_pod')
    result = []
    for book in books:
        link = book.select_one('h3 a')
        result.append({
            'title': link.get('title'),
            'price': soup.select_one('p.price_color').get_text(strip=True),
            'url': link.get('href'),
        })
    return result


def save_to_csv(books, filename):
    if not books:
        print('Нет данных для сохранения')
        return
    with open(filename, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=['title', 'price', 'url'])
        writer.writeheader()
        writer.writerows(books)
    print(f'Сохранено {len(books)} книг в {filename}')


def save_to_xlsx(books, filename):
    if not books:
        print('Нет данных для сохранения')
        return
    wb = Workbook()
    ws = wb.active
    ws.title = filename

    headers = ['title', 'price', 'url']
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_align = Alignment(horizontal='center')
    header_fill = PatternFill(start_color='4472C4', end_color="3a75c4", fill_type='solid')

    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = header_align
        cell.fill = header_fill

    for book in books:
        ws.append([
            book['title'],
            book['price'],
            book['url']
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value)) for cell in col if cell.value is not None)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2

    wb.save(filename)

    print(f'Сохранено {len(books)} книг в {filename}')


def main():
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 (KHTML, like Gecko) "
                      "Chrome/120.0.0.0 Safari/537.36"
    }
    base_url = 'https://books.toscrape.com/catalogue/page-{}.html'

    all_books = []
    
    for page_num in range(1, 4):
        print(f'Парсим страницу {page_num}...')
        url = base_url.format(page_num)
        html = get_page(url, headers)
        if html:
            books = parse_page(html)
            all_books.extend(books)
            print(f'  Найдено: {len(books)} книг')
        time.sleep(2)

    save_to_csv(all_books, 'day13_books.csv')
    save_to_xlsx(all_books, 'day13_books.xlsx')


if __name__ == '__main__':
    main()