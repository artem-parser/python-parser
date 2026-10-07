import requests
from bs4 import BeautifulSoup
import csv
import time
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill



def get_page(url, headers):
    try:
        response = requests.get(url, headers, timeout=10)
        response.encoding = 'utf-8'

        if response.status_code == 200:
            return response.text
        print(f'  Статус {response.status_code}')
        return None
    except Exception as e:
        print(f'  Ошибка: {e}')
        return None


def parse_catalog(html, base_url):
    soup = BeautifulSoup(html, 'html.parser')
    links = []
    for book in soup.select('article.product_pod h3 a'):
        href = book.get('href')
        if href.startswith('catalogue/'):
            full_url = base_url + href
        else:
            full_url = base_url + 'catalogue/' + href
        links.append(full_url)
    return links


def parse_book_cards(html):
    soup = BeautifulSoup(html, 'html.parser')

    title_tag = soup.select_one('h1')
    title = title_tag.get_text(strip=True) if title_tag else 'неизвестно'

    price_tag = soup.select_one('p.price_color')
    price = price_tag.get_text(strip=True) if price_tag else 'неизвестно'

    description_tag = soup.select_one('#product_description ~ p')
    description = description_tag.get_text(strip=True) if description_tag else 'неизвестно'

    stock_tag = soup.select_one('p.instock.availability')
    stock = stock_tag.get_text(strip=True) if stock_tag else 'неизвестно'

    category_tag = soup.select('ul.breadcrumb li a')
    category = category_tag[-1].get_text(strip=True) if category_tag else 'неизвестно'

    img_tag = soup.select_one('#product_gallery img')
    img_url = img_tag.get('src') if img_tag else 'неизвестно'

    return {
        'title': title,
        'price': price,
        'description': description[:200],
        'stock': stock,
        'category': category,
        'img': img_url,
    }


def save_to_csv(books, filename):
    if not books:
        print('Нет данных')
        return
    with open(filename, 'w', newline='', encoding='utf-8-sig') as f:
        fieldnames = ['id', 'title', 'price', 'category', 'description', 'stock', 'img']
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for i, book in enumerate(books, start=1):
            book['id'] = i
            writer.writerow(book)
    print(f'CSV сохранен: {filename}')


def save_to_xlsx(books, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Книги'

    headers = ['id', 'title', 'price',  'category', 'description', 'stock','img']
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for i, book in enumerate(books, start=1):
        ws.append([i, book['title'], book['price'], book['category'],
                   book['description'], book['stock'], book['img']])

    for col in ws.columns:
        max_len = max(len(str(cell.value)) for cell in col if cell.value is not None)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

    wb.save(filename)
    print(f'XLSX сохранен: {filename}')


def main():
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 (KHTML, like Gecko) "
                      "Chrome/120.0.0.0 Safari/537.36"
    }
    base_url = 'https://books.toscrape.com/'
    catalog_url = base_url + 'catalogue/page-{}.html'

    all_books = []
    total_pages = 3

    all_links = []
    for page_num in range(1, total_pages + 1):
        print(f'Каталог: страница {page_num}...')
        html = get_page(catalog_url.format(page_num), headers)
        if html:
            links = parse_catalog(html, base_url)
            all_links.extend(links)
            print(f'  Ссылок собранно: {len(links)}')
        time.sleep(1)

    for i, url in enumerate(all_links, start=1):
        print(f'Карточка {i}/{len(all_links)}: {url[-40:]}')
        html = get_page(url, headers)
        if html:
            book = parse_book_cards(html)
            all_books.append(book)
        time.sleep(1)

    print(f'\nВсего книг: {len(all_books)}')
    print(set(b['category'] for b in all_books))
    save_to_csv(all_books, 'catalog_full.csv')
    save_to_xlsx(all_books, 'catalog_full.xlsx')


if __name__ == '__main__':
    main()