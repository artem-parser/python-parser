# from dataclasses import replace

import requests
from bs4 import  BeautifulSoup
import csv
import time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def get_page(url, headers):
    try:
        response = requests.get(url, headers=headers, timeout=10)
        response.encoding = 'utf-8'

        if response.status_code == 200:
            return response.text
        print(f'Статус: {response.status_code}')
        return None
    except Exception as e:
        print(f'  Ошибка: {e}')
        return None


def parse_page(html):
    soup = BeautifulSoup(html, 'html.parser')
    books = soup.select('article.product_pod')
    result = []

    for book in books:
        link = book.select_one('h3 a')
        prise_text = book.select_one('p.price_color').get_text(strip=True)
        rating_tag = book.select_one('p.star-rating')
        rating = rating_tag['class'][1] if rating_tag else 'нет'
        stock = book.select_one('p.instock.availability')
        stock_text = stock.get_text(strip=True) if stock else 'неизвестно'
        result.append({
            'title': link.get('title'),
            'price': prise_text,
            'rating': rating,
            'stock': stock_text,
            'url': link.get('href')
        })
    return result

def save_to_csv(books, filename):
    with open(filename, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=['id', 'title', 'price', 'rating', 'stock', 'url'])
        writer.writeheader()
        for i, book in enumerate(books, start=1):
            book['id'] = i
            writer.writerow(book)
    print(f'CSV сохранён: {filename}')


def save_to_xlsx(books, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Книги'

    headers = ['id', 'title', 'price', 'rating', 'stock', 'url']
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_align = Alignment(horizontal='center')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for i, book in enumerate(books, start=1):
        ws.append([i, book['title'], book['price'], book['rating'], book['stock'], book['url']])

    for col in ws.columns:
        max_len = max(len(str(cell.value)) for cell in col if cell.value is not None)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2

    wb.save(filename)
    print(f'Excel сохранён: {filename}')


def main():
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 (KHTML, like Gecko) "
                      "Chrome/120.0.0.0 Safari/537.36"
    }
    base_url = 'https://books.toscrape.com/catalogue/page-{}.html'
    total_pages = 5
    all_books = []

    print(f'Начинаем парсинг {total_pages} страниц...')
    for page_num in range(1, total_pages + 1):
        print(f'Страница {page_num}...')
        url = base_url.format(page_num)
        html = get_page(url, headers)
        if html:
            books = parse_page(html)
            all_books.extend(books)
            print(f'  Найдено: {len(books)} книг')
        time.sleep(2)

    print(f'\nВсего книг: {len(all_books)}')
    save_to_csv(all_books, 'final_books.csv')
    save_to_xlsx(all_books, 'final_books.xlsx')

    five_count = sum(1 for b in all_books if b['rating'] == 'Five')
    one_count = sum(1 for b in all_books if b['rating'] == 'One')
    prices = [float(b['price'].replace('£', '')) for b in all_books]
    avg_price = sum(prices) / len(prices)
    print(f'Книг с рейтином Five: {five_count}')
    print(f'Книг с рейтином One: {one_count}')
    print(f'Средняя цена: £{avg_price:.2f}')


if __name__ == '__main__':
    main()
